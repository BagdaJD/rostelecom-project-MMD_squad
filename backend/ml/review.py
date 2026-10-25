"""Обычная и слепая экспертные таблицы без автоматически заполненных оценок."""

import argparse
import hashlib
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from ml.evaluation import evaluate
from ml.needs import UNKNOWN

PROJECT_HEADERS = [
    "project_id",
    "title",
    "description",
    "baseline_recommendations",
    "hybrid_recommendations",
    "expert_explicit_categories",
    "explicit_evidence_json",
    "expert_comment",
    "split",
]
REVIEW_HEADERS = [
    "review_id",
    "project_id",
    "title",
    "description",
    "method",
    "rank",
    "need_category",
    "evidence_type",
    "evidence",
    "reason",
    "expert_relevance",
    "direct_mention",
    "expert_comment",
    "split",
]


def _style(sheet, *, relevance=False):
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="463277")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.row_dimensions[row[0].row].height = 75
    for col in sheet.columns:
        sheet.column_dimensions[col[0].column_letter].width = min(
            max(len(str(c.value or "")) for c in col) + 2, 65
        )
    if relevance:
        for column, options in (("K", "релевантно,нерелевантно,неясно"), ("L", "да,нет,неясно")):
            validation = DataValidation(type="list", formula1=f'"{options}"', allow_blank=True)
            sheet.add_data_validation(validation)
            validation.add(f"{column}2:{column}{sheet.max_row}")


def export_review(baseline: dict, hybrid: dict, output: Path | str, *, seed=20261024):
    output = Path(output)
    output.mkdir(parents=True, exist_ok=True)
    b = {p["external_id"]: p for p in baseline["projects"]}
    h = {p["external_id"]: p for p in hybrid["projects"]}
    if set(b) != set(h):
        raise ValueError("Оба метода должны содержать один и тот же полный набор проектов")
    order = sorted(b, key=lambda pid: hashlib.sha256(f"{seed}:{pid}".encode()).hexdigest())
    split = {pid: "dev" if i < len(order) // 3 else "test" for i, pid in enumerate(order)}
    key = {}
    for blind in (False, True):
        book = Workbook()
        instructions = book.active
        instructions.title = "Instructions"
        instructions.append(["Экспертная проверка: не копируйте рекомендации в истинные метки"])
        instructions.append(
            ["Пустое поле = не размечено. Оценки и прямые упоминания заполняются вручную."]
        )
        instructions.append(
            ["explicit_categories: полный JSON список явных категорий, включая пропуски методов."]
        )
        instructions.append(
            ["[] = проверенное отсутствие; explicit_evidence_json: category -> точная цитата."]
        )
        instructions.append(
            ["Неясно не считается нерелевантным. Dev/test фиксированы до получения оценок."]
        )
        instructions.append(
            ["insufficient_data — отказ; строка оставлена для покрытия всех проектов."]
        )
        projects = book.create_sheet("Projects")
        projects.append(
            [
                *PROJECT_HEADERS[:3],
                "expert_explicit_categories",
                "explicit_evidence_json",
                "expert_comment",
                "split",
            ]
            if blind
            else PROJECT_HEADERS
        )
        recommendations = book.create_sheet("Recommendations")
        headers = REVIEW_HEADERS.copy()
        if blind:
            headers[4] = "variant"
        recommendations.append(headers)
        for pid in sorted(b):
            project = b[pid]
            if (project["title"], project.get("description")) != (
                h[pid]["title"],
                h[pid].get("description"),
            ):
                raise ValueError("Тексты сравниваемых проектов различаются")
            row = [pid, project["title"], project.get("description")]
            if not blind:
                row.extend(
                    ", ".join(r["need_category"] for r in method[pid]["recommendations"])
                    for method in (b, h)
                )
            projects.append([*row, None, None, None, split[pid]])
            unique = {}
            for method in ("baseline", "hybrid"):
                for rank, recommendation in enumerate(
                    (b if method == "baseline" else h)[pid]["recommendations"], 1
                ):
                    identity = json.dumps(
                        {
                            k: recommendation.get(k)
                            for k in ("need_category", "evidence_type", "evidence", "reason")
                        },
                        sort_keys=True,
                    )
                    rid = hashlib.sha256(f"{seed}:{pid}:{identity}".encode()).hexdigest()[:20]
                    item = unique.setdefault(rid, {"recommendation": recommendation, "ranks": {}})
                    item["ranks"][method] = rank
            review_ids = sorted(unique) if blind else list(unique)
            for variant, rid in enumerate(review_ids, 1):
                recommendation, ranks = unique[rid]["recommendation"], unique[rid]["ranks"]
                key[rid] = {"project_id": pid, "methods": list(ranks), "ranks": ranks}
                recommendations.append(
                    [
                        rid,
                        pid,
                        project["title"],
                        project.get("description"),
                        f"Вариант {variant}" if blind else ",".join(ranks),
                        variant if blind else ";".join(f"{m}:{r}" for m, r in ranks.items()),
                        recommendation["need_category"],
                        recommendation.get("evidence_type"),
                        recommendation.get("evidence"),
                        recommendation.get("reason"),
                        None,
                        None,
                        None,
                        split[pid],
                    ]
                )
        _style(projects)
        _style(recommendations, relevance=True)
        for row in instructions:
            row[0].alignment = Alignment(wrap_text=True)
        instructions.column_dimensions["A"].width = 120
        book.save(output / ("blind-review.xlsx" if blind else "expert-review.xlsx"))
    (output / "blind-key.json").write_text(
        json.dumps(key, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return {
        "projects": len(b),
        "dev": sum(v == "dev" for v in split.values()),
        "test": sum(v == "test" for v in split.values()),
        "expert_labels": "empty",
    }


def evaluate_review(workbook, baseline, hybrid, *, blind_key=None):
    book = load_workbook(workbook, read_only=True, data_only=True)

    def rows(sheet):
        values = iter(book[sheet].values)
        headers = next(values)
        return [dict(zip(headers, row, strict=True)) for row in values]

    projects, ratings = rows("Projects"), rows("Recommendations")
    if len({p["project_id"] for p in projects}) != len(projects):
        raise ValueError("Повторный project_id в экспертной таблице")
    if any(p["split"] not in ("dev", "test") for p in projects):
        raise ValueError("Split должен быть dev или test")
    key = json.loads(Path(blind_key).read_text(encoding="utf-8")) if blind_key else {}
    if any("variant" in r for r in ratings) and not key:
        raise ValueError("Для слепой таблицы нужен отдельный blind-key.json")
    grouped = {}
    for rating in ratings:
        if key:
            decoded = key[rating["review_id"]]
            if decoded["project_id"] != rating["project_id"]:
                raise ValueError("Неверный ключ слепой таблицы")
            methods = decoded["methods"]
        else:
            methods = rating["method"].split(",")
        for method in methods:
            if method not in ("baseline", "hybrid"):
                raise ValueError("Неизвестный метод в таблице")
            grouped.setdefault((rating["project_id"], method), []).append(rating)
    report = {}
    for method, data in (("baseline", baseline), ("hybrid", hybrid)):
        predictions = {p["external_id"]: p["recommendations"] for p in data["projects"]}
        if set(predictions) != {p["project_id"] for p in projects}:
            raise ValueError("В таблице должны присутствовать все проекты")
        gold, resolved, irrelevant, unclear = [], 0, 0, 0
        by_split = {s: {"resolved": 0, "irrelevant": 0, "unclear": 0} for s in ("dev", "test")}
        for project in projects:
            candidates = grouped.get((project["project_id"], method), [])
            if len(candidates) != len(predictions[project["project_id"]]):
                raise ValueError("Из таблицы удалены или продублированы рекомендации")
            relevant = {}
            for row in candidates:
                original = [
                    p
                    for p in predictions[project["project_id"]]
                    if all(
                        p.get(k) == row.get(k)
                        for k in ("need_category", "evidence_type", "evidence", "reason")
                    )
                ]
                if not original:
                    raise ValueError(
                        "Содержимое рекомендации было изменено; оцените исходный текст"
                    )
                if row["need_category"] == UNKNOWN or row.get("evidence_type") != "inferred":
                    continue
                value = row.get("expert_relevance")
                if value not in (None, "", "релевантно", "нерелевантно", "неясно"):
                    raise ValueError("Неизвестная оценка эксперта")
                unclear += value == "неясно"
                by_split[project["split"]]["unclear"] += value == "неясно"
                if value in ("релевантно", "нерелевантно"):
                    relevant[row["need_category"]] = value == "релевантно"
                    resolved += 1
                    irrelevant += value == "нерелевантно"
                    by_split[project["split"]]["resolved"] += 1
                    by_split[project["split"]]["irrelevant"] += value == "нерелевантно"
            expected = {
                p["need_category"]
                for p in predictions[project["project_id"]]
                if p.get("evidence_type") == "inferred"
            }
            explicit = project.get("expert_explicit_categories")
            gold.append(
                {
                    "project_id": project["project_id"],
                    "text": "\n".join(
                        t for t in (project["title"], project.get("description")) if t
                    ),
                    "split": project["split"],
                    "explicit_categories": json.loads(explicit) if explicit else None,
                    "explicit_evidence": json.loads(project.get("explicit_evidence_json") or "{}"),
                    "inferred_relevant": relevant
                    if expected and set(relevant) == expected
                    else None,
                }
            )
        report[method] = {
            split: evaluate({"projects": gold}, predictions, split=split)
            for split in ("dev", "test")
        }
        report[method]["rating_coverage"] = {
            "resolved_recommendations": resolved,
            "unclear_recommendations": unclear,
            "irrelevant_share_of_resolved": irrelevant / resolved if resolved else None,
        }
        report[method]["rating_coverage_by_split"] = {
            split: {
                **counts,
                "irrelevant_share_of_resolved": counts["irrelevant"] / counts["resolved"]
                if counts["resolved"]
                else None,
            }
            for split, counts in by_split.items()
        }
    book.close()
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--baseline", type=Path, required=True)
    parser.add_argument("--hybrid", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--completed", type=Path)
    parser.add_argument("--blind-key", type=Path)
    args = parser.parse_args()
    baseline = json.loads(args.baseline.read_text(encoding="utf-8"))
    hybrid = json.loads(args.hybrid.read_text(encoding="utf-8"))
    if args.completed:
        args.output.write_text(
            json.dumps(
                evaluate_review(args.completed, baseline, hybrid, blind_key=args.blind_key),
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )
    else:
        print(json.dumps(export_review(baseline, hybrid, args.output)))


if __name__ == "__main__":
    main()
