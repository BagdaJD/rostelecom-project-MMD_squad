"""Экспорт CSV/CRM и DOCX (модуль exports, раздел 12 ТЗ).

CSV: UTF-8 BOM, RFC 4180, русские заголовки, ISO-8601 для дат.
DOCX: редактируемый шаблон КП с дисклеймером и версией модели.
"""

from __future__ import annotations

import csv
import io
from datetime import UTC, datetime
from pathlib import Path

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.models import Company, Deal, DealEstimate, Industry, ProposalExport, Region, SourceRecord
from modules.investments import service as investments
from modules.radar import service as radar

CSV_HEADERS = {
    "companies": {
        "id": "Идентификатор",
        "name": "Название",
        "inn": "ИНН",
        "client_status": "Статус клиента",
        "industry": "Отрасль",
        "region": "Регион",
        "region_code": "Код региона",
        "revenue_latest": "Выручка, млн руб",
        "revenue_year": "Год выручки",
        "revenue_yoy": "Динамика выручки, %",
        "source_note": "Источник и дата",
        "source_url": "Ссылка на источник",
        "is_synthetic": "Синтетическая компания",
        "revenue_is_synthetic": "Синтетическая финансовая опора",
    }
}
CSV_HEADERS["investments"] = {
    "id": "Идентификатор",
    "name": "Название проекта",
    "inn": "ИНН инициатора",
    "company": "Инициатор",
    "stage": "Стадия",
    "stage_method": "Основание стадии",
    "region": "Регион",
    "region_code": "Код региона",
    "municipality": "Муниципалитет",
    "industry": "Отрасль",
    "start": "Начало",
    "end": "Плановый ввод",
    "budget_rub": "Бюджет, руб.",
    "source_note": "Источник и дата",
    "source_url": "Ссылка на источник",
    "is_synthetic": "Синтетический проект",
}
CSV_HEADERS["deal_estimate"] = {
    "section": "Раздел",
    "name": "Показатель",
    "value": "Значение",
    "note": "Пояснение",
}
CRM_COLUMNS = ("externalId", "name", "inn", "industry", "region", "phone", "website")


def _status_label(status: str) -> str:
    return {"active": "действующий", "prospect": "потенциальный"}.get(status, "не определён")


def companies_rows(session: Session, filters: dict, *, limit=5000, offset=0) -> list[dict]:
    allowed = {
        "region_id",
        "industry_id",
        "industry_group",
        "client_status",
        "min_revenue",
        "min_growth",
        "municipality",
        "q",
        "report_year",
    }
    if set(filters) - allowed:
        raise ValueError("Неподдерживаемые фильтры экспорта компаний")
    companies = radar.list_companies(session, **filters, limit=limit, offset=offset)
    industries = dict(session.execute(select(Industry.id, Industry.name)).all())
    regions = dict(session.execute(select(Region.id, Region.name)).all())
    codes = dict(session.execute(select(Region.id, Region.code)).all())
    sources = {
        source.id: source
        for source in session.scalars(
            select(SourceRecord).where(
                SourceRecord.id.in_([c.source_id for c in companies if c.source_id])
            )
        )
    }
    rows = []
    for c in companies:
        rows.append(
            {
                "id": c.id,
                "name": c.name,
                "inn": c.inn or "",
                "client_status": _status_label(c.client_status),
                "industry": industries.get(c.industry_id or -1, ""),
                "region": regions.get(c.region_id or -1, ""),
                "region_code": codes.get(c.region_id or -1, ""),
                "revenue_latest": c.revenue_latest if c.revenue_latest is not None else "",
                "revenue_year": c.revenue_year or "",
                "revenue_yoy": c.revenue_yoy if c.revenue_yoy is not None else "",
                "is_synthetic": c.is_synthetic,
                "revenue_is_synthetic": c.revenue_is_synthetic,
                "source_note": _source_note(sources.get(c.source_id)),
                "source_url": (sources[c.source_id].raw_payload or {}).get("source_url", "")
                if c.source_id in sources
                else "",
            }
        )
    return rows


def _source_note(source: SourceRecord | None) -> str:
    if source is None:
        return ""
    from modules.radar.companies import provenance

    meta = provenance(source)
    return " | ".join(
        str(value)
        for value in (
            meta.source,
            meta.fetched_at.isoformat(),
            meta.source_file,
            meta.sheet,
            meta.row,
        )
        if value is not None
    )


def investment_rows(session: Session, filters: dict, *, limit=5000, offset=0) -> list[dict]:
    allowed = {"region_id", "industry_id", "stage", "initiator_company_id", "municipality", "q"}
    if set(filters) - allowed:
        raise ValueError("Неподдерживаемые фильтры экспорта инвестпроектов")
    projects = investments.list_projects(session, **filters, limit=limit, offset=offset)
    companies = {
        c.id: c
        for c in session.scalars(
            select(Company).where(
                Company.id.in_([p.initiator_company_id for p in projects if p.initiator_company_id])
            )
        )
    }
    regions = {r.id: r for r in session.scalars(select(Region))}
    industries = dict(session.execute(select(Industry.id, Industry.name)).all())
    sources = {
        s.id: s
        for s in session.scalars(
            select(SourceRecord).where(
                SourceRecord.id.in_([p.source_id for p in projects if p.source_id])
            )
        )
    }
    rows = []
    for p in projects:
        company = companies.get(p.initiator_company_id)
        region = regions.get(p.region_id)
        rows.append(
            {
                "id": p.id,
                "name": p.title,
                "company": company.name if company else "",
                "inn": company.inn if company and company.inn else "",
                "stage": p.stage,
                "stage_method": p.stage_method,
                "is_synthetic": p.is_synthetic,
                "industry": industries.get(p.industry_id, ""),
                "region": region.name if region else "",
                "region_code": region.code if region else "",
                "municipality": p.municipality or "",
                "budget_rub": p.estimated_budget,
                "start": _date_label(p.planned_start, p.start_precision),
                "end": _date_label(p.planned_end, p.end_precision),
                "source_note": _source_note(sources.get(p.source_id)),
                "source_url": p.source_url or "",
            }
        )
    return rows


def _date_label(value, precision):
    if value is None:
        return ""
    return str(value.year) if precision == "year" else value.isoformat()


def _csv_value(value):
    if value is None:
        return ""
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def write_csv(buffer, rows, module: str, mode: str = "general", *, header=True):
    headers = CSV_HEADERS.get(module, {})
    if mode == "crm":
        writer = csv.DictWriter(buffer, fieldnames=list(CRM_COLUMNS), quoting=csv.QUOTE_MINIMAL)
        if header:
            writer.writeheader()
        for row in rows:
            writer.writerow(
                {
                    "externalId": row.get("id", ""),
                    "name": _csv_value(row.get("name", "")),
                    "inn": _csv_value(row.get("inn", "")),
                    "industry": _csv_value(row.get("industry", "")),
                    "region": _csv_value(row.get("region", "")),
                    "phone": "",
                    "website": "",
                }
            )
    else:
        keys = list(headers) if headers else list(rows[0] if rows else [])
        writer = csv.writer(buffer, quoting=csv.QUOTE_MINIMAL)
        if header:
            writer.writerow([headers.get(k, k) for k in keys])
        for row in rows:
            writer.writerow([_csv_value(row.get(k, "")) for k in keys])


def to_csv(rows: list[dict], module: str, mode: str = "general") -> bytes:
    buffer = io.StringIO()
    write_csv(buffer, rows, module, mode)
    return b"\xef\xbb\xbf" + buffer.getvalue().encode("utf-8")


def deal_estimate_rows(session: Session, estimate_id: int) -> list[dict]:
    """CSV старого расчёта (/deals/estimate): вилка, поправки и аналоги одной таблицей."""
    estimate = session.get(DealEstimate, estimate_id)
    if estimate is None:
        raise ValueError("Оценка не найдена")
    rows = [
        {
            "section": "Расчёт",
            "name": "Метод",
            "value": estimate.estimator_mode,
            "note": f"качество данных: {estimate.data_quality}; модель {estimate.model_version}",
        },
        *(
            {"section": "Вилка", "name": label, "value": value, "note": "руб."}
            for label, value in (
                ("Минимум", estimate.price_min),
                ("Оптимум", estimate.price_optimum),
                ("Максимум", estimate.price_max),
            )
        ),
    ]
    if estimate.price_period:
        rows.append({"section": "Расчёт", "name": "Период цен", "value": estimate.price_period})
    for adjustment in estimate.adjustments_json or []:
        rows.append(
            {
                "section": "Поправка",
                "name": adjustment.get("factor"),
                "value": adjustment.get("value"),
                "note": adjustment.get("note"),
            }
        )
    for comparable in estimate.comparables_json or []:
        rows.append(
            {
                "section": "Аналог",
                "name": comparable.get("title"),
                "value": comparable.get("contract_value"),
                "note": (
                    f"сходство {comparable.get('similarity')}; "
                    f"основание цены: {comparable.get('retro_price_basis')}"
                ),
            }
        )
    return rows


def build_proposal_docx(session: Session, estimate_id: int, author_id: int | None) -> str:
    from docx import Document

    estimate = session.get(DealEstimate, estimate_id)
    if estimate is None:
        raise ValueError("Оценка не найдена")
    deal = session.get(Deal, estimate.deal_id)
    params = (deal.input_params_json or {}) if deal else {}

    doc = Document()
    doc.add_heading("Коммерческое предложение (черновик)", level=1)
    doc.add_paragraph(
        f"Сформировано: {datetime.now(UTC):%Y-%m-%d %H:%M} UTC · "
        f"версия модели: {estimate.model_version}"
    )

    doc.add_heading("Параметры проекта", level=2)
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for key, value in params.items():
        row = table.add_row().cells
        row[0].text = key
        row[1].text = ", ".join(map(str, value)) if isinstance(value, list) else str(value)

    doc.add_heading("Оценка стоимости", level=2)
    doc.add_paragraph(f"Минимум: {estimate.price_min:,.0f} руб.")
    doc.add_paragraph(f"Оптимум: {estimate.price_optimum:,.0f} руб.")
    doc.add_paragraph(f"Максимум: {estimate.price_max:,.0f} руб.")
    doc.add_paragraph(f"Метод расчёта: {estimate.estimator_mode}")
    doc.add_paragraph(f"Качество данных: {estimate.data_quality}")
    if estimate.price_period:
        doc.add_paragraph(f"Период ретро-цен: {estimate.price_period}")

    doc.add_heading("Аналоги", level=2)
    for comparable in estimate.comparables_json or []:
        doc.add_paragraph(
            f"{comparable.get('title')} — {comparable.get('contract_value'):,.0f} руб. "
            f"(сходство {comparable.get('similarity')})",
            style="List Bullet",
        )
    if not estimate.comparables_json:
        doc.add_paragraph("Аналоги не подобраны — оценка требует ручной проверки.")

    doc.add_heading("Поправки", level=2)
    for adjustment in estimate.adjustments_json or []:
        doc.add_paragraph(
            f"{adjustment.get('factor')}: {adjustment.get('value')}", style="List Bullet"
        )
    if not estimate.adjustments_json:
        doc.add_paragraph("Поправки не применялись.")

    doc.add_heading("Ограничения", level=2)
    doc.add_paragraph(
        "Документ создан автоматически на основе параметров, введённых пользователем, "
        "и ретро-данных указанного периода. Оценка требует проверки специалистом: "
        "окончательную стоимость сделки утверждает сотрудник."
    )

    out_dir = Path(get_settings().export_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now(UTC).strftime("%Y%m%dT%H%M%SZ")
    path = out_dir / f"proposal-{estimate_id}-{stamp}.docx"
    doc.save(path)

    session.add(
        ProposalExport(
            deal_estimate_id=estimate_id, format="docx", file_ref=str(path), created_by=author_id
        )
    )
    return str(path)
