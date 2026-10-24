"""Коммерческое предложение в DOCX по сохранённому расчёту калькулятора.

Два варианта одного документа:

* ``client`` — для заказчика: что предлагаем, состав, стоимость по статьям, оплата, сроки.
  Себестоимости, маржи и аналогов здесь нет — это коммерческая тайна;
* ``internal`` — то же + приложение «Обоснование цены»: себестоимость, маржа, шанс выиграть,
  поправки с источниками, аналоги, версия модели (требование ТЗ 12 к DOCX).

Цена в КП — утверждённая сотрудником, а если её нет — оптимум расчёта.
"""

from __future__ import annotations

import io
import math
from datetime import UTC, datetime

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from sqlalchemy.orm import Session

from ml.deal_pricing import payment_options
from modules.deal_calculator.models import CalcEstimate

PURPLE = RGBColor(0x77, 0x00, 0xFF)
INK = RGBColor(0x10, 0x18, 0x28)
MUTED = RGBColor(0x5A, 0x64, 0x72)
WARN = RGBColor(0xB3, 0x36, 0x0A)
HEADER_FILL = "F1E8FF"

LEVEL_TEXT = {
    "basic": "мониторинг: сбор показаний и визуализация",
    "advanced": "мониторинг и дистанционное управление оборудованием",
    "expert": "мониторинг и автоматическое регулирование по заданным алгоритмам",
}
SYSTEM_BENEFITS = {
    "Электроснабжение": "Учёт электроэнергии и контроль качества сети в реальном времени.",
    "Отопление": "Учёт тепла и регулирование температуры без ручного обхода.",
    "Водоснабжение": "Учёт воды и мгновенные оповещения о протечках.",
    "Водоотведение": "Контроль насосных и уровней — авария видна до затопления.",
    "Вентиляция": "Контроль качества воздуха (CO₂) и управление приточными установками.",
    "Кондиционирование": "Управление климатом по расписанию и снижение лишнего расхода.",
    "Газоснабжение": "Газоанализ и автоматическое перекрытие подачи при утечке.",
}
FACTOR_LABELS = {
    "project_level": "Уровень проекта",
    "regional_coefficient": "Регион",
    "urgency": "Срочность",
    "unique_solution": "Уникальность",
    "complexity_mode": "Трактовка сложности",
}
CLIENT_ARTICLES = {
    "equipment": "Оборудование и материалы",
    "installation": "Монтаж",
    "design": "Проектирование (ПИР): схемы, ТЗ, кабельные журналы",
    "commissioning": "Пусконаладка (ПНР): программирование контроллеров, настройка датчиков",
    "supervision": "Шеф-монтаж: контроль качества монтажных работ",
    "software": "Программное обеспечение",
    "reserve": "Нестандартные решения",
}


def _rub(value: float | None) -> str:
    return "—" if value is None else f"{round(float(value)):,}".replace(",", " ") + " ₽"


class _Doc:
    """Тонкая обёртка над python-docx: единые стили КП."""

    def __init__(self) -> None:
        self.doc = Document()
        section = self.doc.sections[0]
        section.left_margin = section.right_margin = Cm(2)
        section.top_margin = section.bottom_margin = Cm(1.6)
        style = self.doc.styles["Normal"]
        style.font.name = "Arial"
        style.font.size = Pt(10.5)
        style.font.color.rgb = INK
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")

    def para(self, text: str = "", size: float = 10.5, bold: bool = False, color=INK,
             after: float = 4, italic: bool = False):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(after)
        run = p.add_run(text)
        run.font.size = Pt(size)
        run.bold = bold
        run.italic = italic
        run.font.color.rgb = color
        return p

    def heading(self, text: str) -> None:
        p = self.para(text, 13, True, PURPLE, after=4)
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.keep_with_next = True

    def subheading(self, text: str) -> None:
        p = self.para(text, 11, True, PURPLE, after=2)
        p.paragraph_format.keep_with_next = True

    def bullets(self, items: list[str]) -> None:
        for item in items:
            p = self.doc.add_paragraph(item, style="List Bullet")
            p.paragraph_format.space_after = Pt(2)

    def table(self, rows: list[list[str]], widths_cm: list[float], header: bool = True,
              total: bool = False, right: tuple[int, ...] = ()) -> None:
        t = self.doc.add_table(rows=len(rows), cols=len(rows[0]))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "bottom", "insideH"):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:color"), "D5D7DD")
            borders.append(el)
        t._tbl.tblPr.append(borders)
        for i, row in enumerate(rows):
            for j, value in enumerate(row):
                cell = t.cell(i, j)
                cell.width = Cm(widths_cm[j])
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                if i < len(rows) - 1:
                    p.paragraph_format.keep_with_next = True
                if j in right:
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                run = p.add_run(str(value))
                run.font.size = Pt(10)
                if (header and i == 0) or (total and i == len(rows) - 1):
                    run.bold = True
                    shade = OxmlElement("w:shd")
                    shade.set(qn("w:val"), "clear")
                    shade.set(qn("w:fill"), HEADER_FILL)
                    cell._tc.get_or_add_tcPr().append(shade)
                    if header and i == 0:
                        run.font.color.rgb = PURPLE
        t.autofit = False
        for j, width in enumerate(widths_cm):  # LibreOffice читает ширину колонки, Word — ячейки
            t.columns[j].width = Cm(width)
        for row in t.rows:
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)


def _customer(session: Session, estimate: CalcEstimate):
    """Заказчик и проект из навигатора — только чтение; без них КП всё равно собирается."""
    from app.models import Company, InvestmentProject

    try:
        company = session.get(Company, estimate.company_id) if estimate.company_id else None
        project = (session.get(InvestmentProject, estimate.investment_project_id)
                   if estimate.investment_project_id else None)
    except Exception:  # noqa: BLE001 — таблиц навигатора может не быть в базе калькулятора
        session.rollback()
        return None, None
    return company, project


def build_proposal(session: Session, estimate_id: int, variant: str = "client") -> bytes:
    """DOCX-файл КП по расчёту калькулятора (байты — отдаются сразу, на диск не пишутся)."""
    estimate = session.get(CalcEstimate, estimate_id)
    if estimate is None:
        raise LookupError("Расчёт не найден")
    if not estimate.cost_json:
        raise ValueError("Расчёт не содержит себестоимости: пересчитайте с составом или площадью")
    params = estimate.request_json or {}
    cost = estimate.cost_json
    company, project = _customer(session, estimate)

    price = float(estimate.approved_price or estimate.price_optimum)
    factor = price / float(estimate.cost_total) if estimate.cost_total else 1.0
    now = datetime.now(UTC)
    number = f"КП-{now:%Y}-{estimate.id:04d}"
    synthetic = any(c.get("retro_price_basis") == "synthetic"
                    for c in (estimate.comparables_json or []))

    d = _Doc()
    if synthetic:
        d.para("ПРОТОТИП · РАСЧЁТ НА ДЕМО-ДАННЫХ · НЕ ЯВЛЯЕТСЯ ОФЕРТОЙ", 8.5, True, WARN, 10)
    d.para("Коммерческое предложение", 22, True, PURPLE, 2)
    d.para(f"№ {number} от {now:%d.%m.%Y}", 11, color=MUTED, after=12)

    area = cost.get("area_m2")
    object_line = ", ".join(x for x in [
        cost.get("object_label"),
        f"{round(area):,} м²".replace(",", " ") if area else None,
        cost.get("region_label"),
    ] if x)
    customer = company.name if company else "[наименование заказчика]"
    if company and company.inn:
        customer += f", ИНН {company.inn}"
    d.table([
        ["Заказчик", customer],
        ["Объект", object_line or "[объект]"],
        *([["Проект", project.title]] if project else []),
        ["Исполнитель", "ПАО «Ростелеком», коммерческий блок"],
        ["Менеджер", "[ФИО], [телефон], [e-mail]"],
    ], [4, 13], header=False)

    level = params.get("project_level", "basic")
    systems = cost.get("systems") or []
    d.heading("1. Что предлагаем")
    d.para("Создание системы диспетчеризации инженерных систем объекта: сбор показаний "
           "с приборов учёта и датчиков на единую панель, " + LEVEL_TEXT.get(level, "") + ".")
    d.bullets([text for group, text in SYSTEM_BENEFITS.items() if group in systems])

    d.heading("2. Состав системы")
    rows = [["Система", "Оборудование", "Кол-во"]]
    order = {g: i for i, g in enumerate(SYSTEM_BENEFITS)}
    lines = sorted(cost.get("lines", []),
                   key=lambda x: (order.get(x.get("group"), 99), x.get("code") or "~"))
    for line in lines:
        rows.append([line.get("group") or "", line["name"], f"{line['qty']:g} шт."])
    rows.append(["", "Итого приборов и контроллеров", f"{cost.get('points', 0):g} шт."])
    d.table(rows, [4.2, 10.3, 2.5], total=True, right=(2,))

    d.heading("3. Стоимость")
    parts = [(b["article"], float(b["amount"])) for b in cost.get("breakdown", [])]
    sale = [round(v * factor) for _, v in parts]
    if sale:
        sale[0] += round(price) - sum(sale)  # копейки округления — в первую статью
    rows = [["Статья", "Сумма"]]
    rows += [[CLIENT_ARTICLES.get(a, a), _rub(v)] for (a, _), v in zip(parts, sale, strict=True)]
    rows.append(["Итого", _rub(price)])
    d.table(rows, [13, 4], total=True, right=(1,))
    d.para("Цены указаны без НДС. Налог начисляется по ставке, действующей на дату "
           "заключения договора.", 9, color=MUTED, italic=True)

    d.heading("4. Варианты оплаты")
    options = payment_options(price)
    rows = [["Схема", "Платёж", "Итого за период"]]
    for opt in options:
        unit = "" if opt["scheme"] == "capex" else " / мес"
        rows.append([opt["label"], _rub(opt["monthly"]) + unit, _rub(opt["total"])])
    d.table(rows, [8.6, 4.2, 4.2], right=(1, 2))
    chosen = next((o for o in options if o["scheme"] == params.get("payment_scheme")), None)
    if chosen and chosen["scheme"] != "capex":
        d.para(f"Рекомендуемая схема: {chosen['label']}.", 10)

    months = params.get("duration_months") or cost.get("norm_months") or 3
    weeks = max(4, math.ceil(float(months) * 4.3))
    d.heading("5. Сроки")
    d.table([
        ["Этап", "Срок"],
        ["Обследование и проектирование", f"{max(1, round(weeks * 0.25))} нед."],
        ["Поставка оборудования", f"{max(1, round(weeks * 0.35))} нед. (параллельно с проектом)"],
        ["Монтаж, пусконаладка, обучение персонала", f"{max(1, round(weeks * 0.3))} нед."],
        ["Сдача системы", f"{max(1, round(weeks * 0.1))} нед."],
        ["Всего", f"около {float(months):g} мес."],
    ], [10, 7], total=True)

    d.heading("6. Условия")
    d.bullets(["Предложение действительно 30 календарных дней с даты выпуска.",
               "Окончательная стоимость фиксируется в договоре после обследования объекта.",
               "Гарантийные обязательства и уровень сервиса — по условиям договора."])

    if variant == "internal":
        _internal_appendix(d, estimate, price)

    approved = "цена утверждена" if estimate.approved_price else "цена не утверждена"
    d.para("", after=6)
    d.para(f"Сформировано калькулятором сделок {now:%d.%m.%Y %H:%M} UTC · расчёт №{estimate.id}"
           f" · модель {estimate.model_version} · {approved}", 8.5, color=MUTED, after=0)

    buffer = io.BytesIO()
    d.doc.save(buffer)
    return buffer.getvalue()


def _internal_appendix(d: _Doc, estimate: CalcEstimate, price: float) -> None:
    d.doc.add_page_break()
    d.para("ДЛЯ ВНУТРЕННЕГО ПОЛЬЗОВАНИЯ — НЕ ПЕРЕДАВАТЬ ЗАКАЗЧИКУ", 9, True, WARN, 6)
    d.heading("Приложение. Обоснование цены")
    cost_total = float(estimate.cost_total or 0)
    margin = (price - cost_total) / price * 100 if price else 0
    d.table([
        ["Показатель", "Значение"],
        ["Себестоимость", _rub(cost_total)],
        ["Минимум / оптимум / максимум",
         f"{_rub(estimate.price_min)} / {_rub(estimate.price_optimum)} / "
         f"{_rub(estimate.price_max)}"],
        ["Цена в КП", _rub(price)],
        ["Маржа при цене КП", f"{margin:.1f}%".replace(".", ",")],
        ["Шанс выиграть при оптимуме", f"{float(estimate.win_probability):.0f}%"
         if estimate.win_probability is not None else "нет модели"],
        ["Метод", f"{estimate.estimator_mode} · {estimate.optimum_method or '—'}"],
        ["Качество данных", estimate.data_quality],
    ], [7, 10])
    d.subheading("Поправки и их источник")
    d.table([["Поправка", "Значение", "Основание"]] + [
        [FACTOR_LABELS.get(a.get("factor", ""), a.get("factor", "")), f"×{a.get('value')}",
         f"{a.get('note', '')} ({a.get('source', '')})"]
        for a in (estimate.adjustments_json or [])
    ], [4, 2.5, 10.5])
    d.subheading("Аналоги")
    rows = [["Сделка", "Сходство", "Цена", "Наценка", "Итог"]]
    for c in (estimate.comparables_json or [])[:10]:
        rows.append([c.get("title", ""), f"{c.get('similarity', 0) * 100:.0f}%",
                     _rub(c.get("contract_value")),
                     f"×{c.get('price_index'):.2f}" if c.get("price_index") else "—",
                     {"won": "выиграли", "lost": "проиграли"}.get(c.get("outcome"), "—")])
    d.table(rows, [7.0, 2.5, 3, 2, 2.5], right=(1, 2, 3))
