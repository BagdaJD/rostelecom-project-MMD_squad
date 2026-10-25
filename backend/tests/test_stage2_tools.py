import json
from copy import deepcopy
from datetime import date

import pytest
from openpyxl import load_workbook

from ingestion.fns_audit import audit_fns
from ingestion.matching import InitiatorReviewIndex, match_company, match_initiator
from ingestion.normalize import RawCompany
from ingestion.pipeline import analyze
from ingestion.sources.fns_source import LocalFnsSource
from ingestion.validate_combined import validate_combined
from ml.diagnostics import distribution, scoring_diagnostics
from ml.needs import classify_need
from ml.review import evaluate_review, export_review
from tests.test_pipeline_normalization import FNS_XML
from tests.test_pipeline_normalization import tiny_excel as tiny_excel


def corpus():
    projects = [
        {"external_id": "p1", "title": "Завод", "description": None},
        {"external_id": "p2", "title": "Канцелярские товары", "description": None},
    ]
    for p in projects:
        p["recommendations"] = [n.as_dict() for n in classify_need(p["title"])]
    return {"projects": projects}


def test_review_covers_abstentions_deduplicates_and_leaves_labels_empty(tmp_path):
    data = corpus()
    result = export_review(data, deepcopy(data), tmp_path)
    assert result["projects"] == 2 and result["expert_labels"] == "empty"
    for name in ("expert-review.xlsx", "blind-review.xlsx"):
        book = load_workbook(tmp_path / name)
        assert book["Projects"].max_row == 3
        assert book["Recommendations"].max_row == 5  # 3 гипотезы + один отказ, без дублей методов.
        for row in book["Recommendations"].iter_rows(min_row=2):
            assert all(row[i].value is None for i in (10, 11, 12))
        if name.startswith("blind"):
            texts = [str(c.value) for sheet in book for row in sheet for c in row]
            assert not any("baseline" in t or "hybrid" in t for t in texts)
        book.close()
    report = evaluate_review(tmp_path / "expert-review.xlsx", data, data)
    assert report["baseline"]["test"]["status"] == "not_evaluated"
    report = evaluate_review(
        tmp_path / "blind-review.xlsx", data, data, blind_key=tmp_path / "blind-key.json"
    )
    assert report["hybrid"]["test"]["status"] == "not_evaluated"


def test_review_only_uses_manual_complete_ratings_and_rejects_tampering(tmp_path):
    data = corpus()
    export_review(data, data, tmp_path)
    path = tmp_path / "expert-review.xlsx"
    book = load_workbook(path)
    for row in book["Recommendations"].iter_rows(min_row=2):
        if row[6].value != "insufficient_data":
            row[10].value = "релевантно"  # Только синтетическая тестовая разметка.
    book.save(path)
    result = evaluate_review(path, data, data)
    assert result["baseline"]["test"]["inferred_precision_at_3"] == 1
    assert result["baseline"]["test"]["macro_f1"] is None
    book["Recommendations"]["G2"] = "altered"
    book.save(path)
    with pytest.raises(ValueError, match="изменено"):
        evaluate_review(path, data, data)


def test_review_preserves_source_as_text_not_formula(tmp_path):
    data = corpus()
    data["projects"][0]["title"] = '=HYPERLINK("https://example.invalid")'
    export_review(data, data, tmp_path)
    book = load_workbook(tmp_path / "expert-review.xlsx")
    assert book["Projects"]["B2"].data_type == "s"
    assert book["Projects"]["B2"].value.startswith("=")


def test_review_rejects_missing_project_and_blind_without_key(tmp_path):
    data = corpus()
    export_review(data, data, tmp_path)
    with pytest.raises(ValueError, match="blind-key"):
        evaluate_review(tmp_path / "blind-review.xlsx", data, data)
    book = load_workbook(tmp_path / "expert-review.xlsx")
    book["Projects"].delete_rows(2)
    book.save(tmp_path / "expert-review.xlsx")
    with pytest.raises(ValueError, match="все проекты"):
        evaluate_review(tmp_path / "expert-review.xlsx", data, data)


def test_fuzzy_and_geography_candidates_never_confirm_identity():
    c = RawCompany(
        source="test",
        external_id="1",
        name="ООО Тестовый Альфа",
        inn="1234567894",
        region_code="61",
        municipality="Аксайский район",
    )
    project = {
        "initiator": "Тестовый Альфa",
        "region_code": "61",
        "municipality": "г. Ростов-на-дону",
    }
    decision = match_initiator(project, [c], review_index=InitiatorReviewIndex([c]))
    assert decision.status == "requires_review" and decision.confidence < 1
    assert decision.candidate_indices == (0,)
    assert match_company(c, [c, deepcopy(c)]).status == "requires_review"


def test_fns_rejects_invalid_checksums_without_logging_identity(tmp_path):
    path = tmp_path / "bad-check.xml"
    path.write_text(FNS_XML.replace('ИННЮЛ="1234567894"', 'ИННЮЛ="1234567895"'), encoding="utf-8")
    source = LocalFnsSource(path)
    rows = list(source.iter_companies())
    assert len(rows) == 1 and source.counts["invalid_inn"] == 1
    assert all("1234567895" not in str(e) for e in source.errors)


def test_fns_audit_marks_partial_and_full_scopes(tiny_excel, tmp_path):
    path = tmp_path / "fns.xml"
    path.write_text(FNS_XML, encoding="utf-8")
    partial = audit_fns(path, tiny_excel, tmp_path / "partial", limit=1)
    full = audit_fns(path, tiny_excel, tmp_path / "full")
    assert partial["scope"] == "technical_partial" and not partial["complete"]
    assert full["scope"] == "full_archive" and full["complete"]
    assert full["unique_region_companies"] == 2 and full["excel_matches"] == 1
    assert full["counts"]["invalid_records"] == 1
    assert full["runtime"]["sampled_peak_rss_bytes"] > 0


def test_diagnostics_do_not_change_weights_or_leak_company_names():
    record = RawCompany(
        source="test",
        external_id="1",
        name="Приватное тестовое имя",
        inn="1234567894",
        industry_code="33.12",
        sme_category="Микропредприятие",
    )
    data = analyze([record], [], as_of=date(2026, 10, 24))
    before = json.dumps(data, ensure_ascii=False)
    result = scoring_diagnostics(data)
    assert result["business_weights_changed"] is False
    assert "Приватное тестовое имя" not in json.dumps(result, ensure_ascii=False)
    assert json.dumps(data, ensure_ascii=False) == before
    assert distribution([]) == {"count": 0}


def test_combined_requires_full_manifest_and_keeps_case_status(tiny_excel, tmp_path):
    path = tmp_path / "fns.xml"
    path.write_text(FNS_XML, encoding="utf-8")
    output = tmp_path / "fns-audit"
    audit_fns(path, tiny_excel, output)
    result = validate_combined(
        tiny_excel,
        output / "companies.jsonl",
        output / "quality.json",
        tmp_path / "combined",
        as_of=date(2026, 10, 24),
    )
    assert result["excel_companies"] == 1 and result["fns_companies"] == 2
    assert result["combined_unique_companies"] == 2
    assert result["territories"]["case_clients"] == 1
    assert result["territories"]["other_candidates"] == 1
    manifest = json.loads((output / "quality.json").read_text(encoding="utf-8"))
    manifest["complete"] = False
    (output / "quality.json").write_text(json.dumps(manifest), encoding="utf-8")
    with pytest.raises(ValueError, match="завершённый"):
        validate_combined(
            tiny_excel,
            output / "companies.jsonl",
            output / "quality.json",
            tmp_path / "refused",
            as_of=date(2026, 10, 24),
        )
