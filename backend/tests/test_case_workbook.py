"""The supplied Excel, source updates and the connected HTTP business workflow."""

import csv
import io
import json
from decimal import Decimal

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.db.session import get_session
from app.main import app
from app.models import (
    Company,
    CompanyMatch,
    CompanyMetric,
    InvestmentEvent,
    InvestmentProject,
    ProjectMilestone,
    ProjectNeed,
    SourceRecord,
    User,
)
from ingestion.case_workbook import COMPANY_SOURCE, import_case, read_case
from modules.deal_calculator.models import CalcRegion
from modules.deal_calculator.seed import DEFAULT_BUNDLE, load_bundle


def counts(session):
    return [
        session.scalar(select(func.count()).select_from(model))
        for model in (
            Company,
            CompanyMatch,
            CompanyMetric,
            InvestmentProject,
            InvestmentEvent,
            ProjectMilestone,
            ProjectNeed,
            SourceRecord,
        )
    ]


def test_actual_workbook_is_complete_repeatable_and_honest(business_db):
    session, _ = business_db
    path = get_settings().case_workbook_path
    parsed = read_case(path)
    assert len(parsed.companies) == 759 and len(parsed.projects) == 63
    assert parsed.issues == []
    first = import_case(session, path)
    session.commit()
    assert first["companies_created"] == 759
    assert first["initiators_created"] == 60
    assert first["projects_created"] == 63
    before = counts(session)
    second = import_case(session, path)
    session.commit()
    assert second["companies_created"] == second["projects_created"] == 0
    assert second["initiators_created"] == second["projects_updated"] == 0
    assert counts(session) == before
    assert session.scalar(select(func.count(CompanyMetric.id))) == 0
    assert session.scalar(select(func.count(Company.id)).where(Company.employee_count == 0)) > 0
    assert session.scalar(select(func.count(Company.id)).where(Company.is_synthetic)) == 0
    assert (
        session.scalar(select(func.count(Company.id)).where(Company.client_status == "active")) == 0
    )
    assert session.scalar(select(func.count(Company.id)).where(Company.inn.is_(None))) == 60
    project = session.scalar(select(InvestmentProject).order_by(InvestmentProject.id))
    row = parsed.projects[0]
    assert project.estimated_budget == Decimal(row["budget_rub"])
    assert project.start_precision == "year" and project.planned_start.year == row["start_year"]
    source = session.get(SourceRecord, project.source_id)
    assert source.raw_payload["provenance"]["file"] == path.name
    assert source.license_note and source.payload_hash


def test_changed_source_updates_same_objects_and_retains_history(business_db, tmp_path):
    session, _ = business_db
    path = tmp_path / "case.xlsx"
    workbook = load_workbook(get_settings().case_workbook_path)
    workbook.save(path)
    import_case(session, path)
    session.commit()
    company_source = session.scalar(
        select(SourceRecord).where(SourceRecord.source == COMPANY_SOURCE).order_by(SourceRecord.id)
    )
    company = session.scalar(select(Company).where(Company.source_id == company_source.id))
    project = session.scalar(select(InvestmentProject).order_by(InvestmentProject.id))
    old_source, company_id, project_id = company.source_id, company.id, project.id
    company_sheet = next(s for s in workbook if any(c.value == "Наименование / ФИО" for c in s[1]))
    company_row = company_source.raw_payload["attributes"]["provenance"]["row"]
    name_column = next(c.column for c in company_sheet[1] if c.value == "Наименование / ФИО")
    employees = next(
        c.column for c in company_sheet[1] if str(c.value).startswith("Среднесписочная")
    )
    company_sheet.cell(company_row, employees, 0)

    company_sheet.cell(company_row, name_column, company.name + " UPDATED")
    project_sheet = next(
        s for s in workbook if any(c.value == "Инициатор инвестиционного проекта" for c in s[1])
    )
    project_sheet.cell(2, 6, project.title + " UPDATED")
    project_sheet.cell(2, 9, 12.5)
    workbook.save(path)
    workbook.close()
    report = import_case(session, path)
    session.commit()
    session.expire_all()
    assert report["companies_created"] == report["projects_created"] == 0
    assert report["projects_updated"] == 1
    assert session.get(Company, company_id).employee_count == 0
    assert session.get(Company, company_id).source_id != old_source
    assert session.get(SourceRecord, old_source) is not None
    assert session.get(InvestmentProject, project_id).estimated_budget == Decimal("12500000")
    before = counts(session)
    import_case(session, path)
    session.commit()
    assert counts(session) == before


@pytest.fixture
def case_api(business_db, authenticated_session):
    session, engine = business_db
    import_case(session, get_settings().case_workbook_path)
    session.commit()
    load_bundle(session, json.loads(DEFAULT_BUNDLE.read_text(encoding="utf-8")))
    user = User(email="case@example.test", login="case", password_hash="test", role="manager")
    session.add(user)
    session.commit()
    token = authenticated_session(session, user).access_token

    def database():
        with Session(engine) as request_session:
            yield request_session

    app.dependency_overrides[get_session] = database
    try:
        with TestClient(app) as client:
            client.headers["Authorization"] = f"Bearer {token}"
            yield client, session
    finally:
        app.dependency_overrides.pop(get_session, None)


def test_case_cards_reference_data_and_missing_values(case_api):
    client, _ = case_api
    regions = client.get("/api/v1/regions").json()
    assert len(regions) == 1 and regions[0]["code"] == "RU-ROS-RND"
    region_id = regions[0]["id"]
    assert client.get(f"/api/v1/regions/{region_id}/municipalities").json()
    assert len(client.get("/api/v1/industries").json()) == 13
    stages = client.get("/api/v1/dictionaries/project_stage").json()
    assert any(row["code"] == "unknown" for row in stages)
    companies = client.get("/api/v1/companies", params={"limit": 1})
    assert companies.status_code == 200, companies.text
    row = companies.json()[0]
    assert row["revenue_latest"] is None and row["revenue_yoy"] is None
    card = client.get(f"/api/v1/companies/{row['id']}")
    assert card.status_code == 200, card.text
    assert card.json()["metrics"] == [] and card.json()["provenance"]
    source = client.get(f"/api/v1/sources/{card.json()['provenance'][0]['id']}")
    assert source.status_code == 200 and source.json()["raw_payload"]
    assert client.get("/api/v1/companies/999999").status_code == 404
    assert client.get("/api/v1/investments/999999/needs").status_code == 404
    heat = client.get("/api/v1/radar/heatmap", params={"level": "municipality"})
    assert heat.status_code == 200, heat.text
    assert sum(c["companies_count"] for c in heat.json()) == 819
    assert all(c["revenue_sum"] is None and c["presence_score"] is None for c in heat.json())
    assert all(c["longitude"] is None and c["latitude"] is None for c in heat.json())
    assert client.get("/api/v1/radar/white-spaces").json() == []
    timeline = client.get("/api/v1/investments/timeline", params={"period": "month"}).json()
    assert timeline and all(
        len(p["period"]) == 4 and p["date_precision"] == "year" for p in timeline
    )


def test_project_to_estimate_approval_docx_and_filtered_csv(case_api):
    client, session = case_api
    listed = client.get("/api/v1/investments", params={"limit": 500})
    assert listed.status_code == 200 and len(listed.json()) == 63
    project = listed.json()[0]
    project_id = project["id"]
    card = client.get(f"/api/v1/investments/{project_id}")
    assert card.status_code == 200, card.text
    assert card.json()["planned_start"] is None and card.json()["planned_end"] is None
    assert card.json()["start_precision"] == "year" and card.json()["provenance"]
    assert all(m["date"] is None for m in card.json()["milestones"])
    draft = client.get(f"/api/v1/calculator/draft/{project_id}")
    assert draft.status_code == 200, draft.text
    request = draft.json()["request"]
    assert request["investment_project_id"] == project_id
    assert request["company_id"] == project["initiator_company_id"]
    assert request["duration_months"] is None
    assert session.get(CalcRegion, request["region_id"]).code == "RU-ROS-RND"
    estimate = client.post("/api/v1/calculator/estimate", json=request)
    assert estimate.status_code == 200, estimate.text
    result = estimate.json()
    assert Decimal(result["estimate"]["optimum"]) > 0
    estimate_id = result["estimate_id"]
    approve = client.post(
        f"/api/v1/calculator/estimates/{estimate_id}/approve",
        json={"price": float(result["estimate"]["optimum"])},
    )
    assert approve.status_code == 200, approve.text
    docx = client.get(f"/api/v1/calculator/estimates/{estimate_id}/proposal")
    assert docx.status_code == 200, docx.text
    assert Document(io.BytesIO(docx.content)).paragraphs
    for module, filters in (
        ("investments", {"initiator_company_id": request["company_id"]}),
        ("companies", {"q": "an-unmatched-company"}),
    ):
        for mode in ("general", "crm"):
            response = client.post(
                "/api/v1/exports/csv",
                json={
                    "module": module,
                    "mode": mode,
                    "filters": filters,
                },
            )
            assert response.status_code == 200, response.text
            assert response.content.startswith(b"\xef\xbb\xbf")
            rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
            assert len(rows) > 1 if module == "investments" else len(rows) == 1
    invalid = client.post(
        "/api/v1/exports/csv",
        json={
            "module": "investments",
            "filters": {"limit": 9999},
        },
    )
    assert invalid.status_code == 422
    invalid = client.post(
        "/api/v1/exports/csv",
        json={
            "module": "companies",
            "filters": {"min_revenue": "not-a-number"},
        },
    )
    assert invalid.status_code == 422
