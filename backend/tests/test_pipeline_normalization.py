"""Регрессии подготовки данных; все сущности в фикстурах синтетические."""

from datetime import UTC, datetime, timedelta
from zipfile import ZipFile

import openpyxl
import pytest

from ingestion.geography import municipality, normalize_municipality
from ingestion.matching import match_company, match_initiator, unique_companies
from ingestion.normalize import (
    RawCompany,
    clean_text,
    identifier,
    integer,
    is_valid_inn,
    is_valid_ogrn,
    iso_date,
    normalize_name,
    number,
)
from ingestion.sources.excel_source import read_excel
from ingestion.sources.fns_source import LocalFnsSource


def company(**kwargs):
    return RawCompany(
        **{
            "source": "test",
            "external_id": "1",
            "name": "ООО Тестовый объект",
            "region_code": "61",
            "municipality": "г. Ростов-на-дону",
            **kwargs,
        }
    )


@pytest.mark.parametrize("inn", ["1234567894", "012345678943", "500100732259"])
def test_both_inn_checks_and_leading_zero(inn):
    assert is_valid_inn(inn)
    assert not is_valid_inn(inn[:-1] + str((int(inn[-1]) + 1) % 10))


def test_second_individual_control_digit_is_checked():
    assert not is_valid_inn("012345678944")
    assert not is_valid_inn("012345678953")
    assert not is_valid_inn("000000000000")
    assert not is_valid_inn("１２３４５６７８９４")
    assert not is_valid_inn(1234567894)


def test_ogrn_checksums():
    assert is_valid_ogrn("1027700000008")
    assert not is_valid_ogrn("1027700000009")
    assert not is_valid_ogrn("0000000000000")


def test_normalization_missing_and_placeholder_scope():
    assert identifier("012345678943") == "012345678943"
    assert identifier("1234567894.0") == "1234567894"
    assert clean_text("136") == "136"
    assert clean_text("136", placeholder_136=True) is None
    assert number(136) == 136
    assert number("1 250,5") == 1250.5
    assert number("nan") is None and number("inf") is None
    assert number("неизвестно") is None
    assert integer("0") == 0 and integer("2.5") is None
    assert iso_date("10.07.2022") == "2022-07-10"
    assert iso_date("2026") is None
    assert normalize_name("ОБЩЕСТВО С ОГРАНИЧЕННОЙ ОТВЕТСТВЕННОСТЬЮ «Ёлка»") == "елка"
    assert normalize_name("ООО Сервис АО") == "сервис ао"


def test_municipal_variants_and_unknown_region():
    assert municipality("АКСАЙСКИЙ Р-Н") == "Аксайский район"
    assert municipality(city="ГОРОД АКСАЙ") == "Аксайский район"
    assert normalize_municipality("г. Ростов-на-Дону") == "г. Ростов-на-дону"
    assert municipality(city="Неизвестный город") is None


def test_different_valid_inns_never_merge_on_name():
    a = company(inn="1234567894")
    b = company(inn="012345678943", external_id="2")
    decision = match_company(a, [b])
    assert decision.status == "requires_review"
    result, counts = unique_companies([a, b])
    assert len(result) == 2 and counts["duplicates"] == 0


def test_conflicting_identifier_requires_review():
    a = company(inn="1234567894", ogrn="1027700000008")
    b = company(inn="012345678943", ogrn="1027700000008")
    assert match_company(a, [b]).status == "requires_review"


def test_id_dedup_and_case_status_precedence():
    a = company(inn="1234567894", source="case_excel", client_status="active")
    b = company(inn="1234567894", source="fns_rsmp", external_id="fns:1")
    result, counts = unique_companies([b, a])
    assert len(result) == 1 and result[0].client_status == "active"
    assert counts["identifier_matches"] == 1


def test_project_name_match_is_only_candidate():
    result = match_initiator(
        {
            "initiator": "ООО Тестовый объект",
            "region_code": "61",
            "municipality": "г. Ростов-на-дону",
        },
        [company()],
    )
    assert result.status == "requires_review" and result.confidence < 1
    result = match_initiator({"initiator_inn": "1234567894"}, [company(inn="1234567894")])
    assert result.status == "matched"


def test_stable_payload_excludes_fetch_time():
    a = company()
    b = company(fetched_at=datetime.now(UTC) + timedelta(days=1))
    assert a.stable_payload() == b.stable_payload()


@pytest.fixture
def tiny_excel(tmp_path):
    book = openpyxl.Workbook()
    cs = book.active
    cs.title = "clients"
    cs.append(
        [
            "№ п/п",
            "Наименование / ФИО",
            "Тип субъекта",
            "ОГРН",
            "ИНН",
            "Регион",
            "Город",
            "Основной вид деятельности",
            "Категория",
        ]
    )
    cs.append(
        [
            1,
            "ООО Тестовый объект",
            "Юридическое лицо",
            "1027700000008",
            "012345678943",
            "61 - Ростовская область",
            "г. Ростов-на-Дону",
            "33.12 Ремонт",
            "Микропредприятие",
        ]
    )
    ps = book.create_sheet("renamed_projects")
    ps.append(
        [
            "№ п/п",
            "Муниципальное образование",
            "Инициатор инвестиционного проекта",
            "Название и суть проекта",
            "Начало",
            "Ввод в эксплуатацию",
            "Всего (млн руб.)",
            "Количество новых рабочих мест",
        ]
    )
    ps.append(
        [
            1,
            "г. Ростов-на-Дону",
            "ООО Тестовый объект",
            "Строительство завода",
            "2024",
            "2026",
            136,
            0,
        ]
    )
    other = book.create_sheet("МОД3_Таб1")
    other.append(["Лист калькулятора не входит в ML"])
    path = tmp_path / "test.xlsx"
    book.save(path)
    return path


def test_excel_header_detection_and_original_unchanged(tiny_excel):
    original = tiny_excel.read_bytes()
    result = read_excel(tiny_excel)
    assert len(result.companies) == len(result.projects) == 1
    assert result.companies[0].inn == "012345678943"
    assert result.projects[0]["investment_mln_rub"] == 136
    assert result.projects[0]["new_jobs"] == 0
    assert result.projects[0]["stage"] == "unknown"
    assert result.projects[0]["planned_start"] is None
    assert tiny_excel.read_bytes() == original
    assert not any(k.startswith("МОД3") for k in result.quality["sheets"])


FNS_XML = """<?xml version="1.0" encoding="UTF-8"?>
<Файл xmlns="urn:test"><Документ ИдДок="T1" ДатаСост="10.10.2026" ДатаВклМСП="10.01.2020"
 КатСубМСП="1" ССЧР="136"><ОргВклМСП НаимОрг="ООО Тестовый объект" ИННЮЛ="1234567894"
 ОГРН="1027700000008"/><СведМН КодРегион="61"><Город Тип="г" Наим="Ростов-на-Дону"/>
 </СведМН><СвОКВЭД><СвОКВЭДОсн КодОКВЭД="33.12" НаимОКВЭД="Ремонт"/></СвОКВЭД>
 </Документ><Документ ИдДок="T2"><ОргВклМСП НаимОрг="Другой регион" ИННЮЛ="1234567894"/>
 <СведМН КодРегион="77"/></Документ><Документ ИдДок="T3"><СведМН КодРегион="61"/></Документ>
 <Документ ИдДок="T4" КатСубМСП="2"><ИПВклМСП ИННФЛ="012345678943" ОГРНИП="">
 <ФИОИП Фамилия="Тестовый" Имя="Персонаж"/></ИПВклМСП><СведМН КодРегион="61"/>
 </Документ></Файл>"""


@pytest.mark.parametrize("zipped", [False, True])
def test_fns_xml_and_zip_streaming(tmp_path, zipped):
    path = tmp_path / ("test.zip" if zipped else "test.xml")
    if zipped:
        with ZipFile(path, "w") as archive:
            archive.writestr("../input.xml", FNS_XML.encode())
    else:
        path.write_text(FNS_XML, encoding="utf-8")
    source = LocalFnsSource(path)
    rows = list(source.iter_companies())
    assert len(rows) == 2
    assert rows[0].region_code == "61"  # ИНН начинается на 12, не на 61.
    assert rows[0].employee_count == 136
    assert rows[0].dump_date == "2026-10-10"
    assert rows[0].industry_code == "33.12"
    assert rows[0].client_status == "unknown"
    assert source.counts["invalid_records"] == 1
    assert source.counts["other_or_missing_region"] == 1
    assert "Тестовый" not in str(source.errors)
    assert not (tmp_path.parent / "input.xml").exists()


def test_malformed_xml_is_reported_without_raw_data(tmp_path):
    path = tmp_path / "bad.xml"
    path.write_text("<root><unclosed></root>", encoding="utf-8")
    source = LocalFnsSource(path)
    assert list(source.iter_companies()) == []
    assert source.errors[0]["code"] == "malformed_xml"
