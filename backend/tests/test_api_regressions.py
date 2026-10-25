"""HTTP regressions on portable tables; these do not replace PostGIS integration tests."""

import io
import json

import pytest
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.core.config import get_settings
from app.db.base import Base
from app.db.session import get_session
from app.main import app
from app.models import (
    AuditLog,
    Deal,
    DealBreakdown,
    DealEstimate,
    HistoricalDeal,
    Notification,
    PositionPrice,
    ProposalExport,
    RefreshToken,
    User,
)
from modules.deal_calculator.seed import DEFAULT_BUNDLE, create_tables
from modules.deal_calculator.seed import load_bundle as load_calculator_bundle
from modules.deals.service import FALLBACK_MODE


@pytest.fixture
def api_client(tmp_path, monkeypatch, authenticated_session):
    engine = create_engine(
        "sqlite://", poolclass=StaticPool, connect_args={"check_same_thread": False}
    )
    tables = (
        User,
        RefreshToken,
        AuditLog,
        Notification,
        HistoricalDeal,
        PositionPrice,
        Deal,
        DealEstimate,
        DealBreakdown,
        ProposalExport,
    )
    Base.metadata.create_all(engine, tables=[model.__table__ for model in tables])
    with Session(engine) as session:
        user = User(
            email="owner@example.test", login="owner", password_hash="test-only", role="manager"
        )
        other = User(
            email="other@example.test", login="other", password_hash="test-only", role="manager"
        )
        session.add_all([user, other])
        session.commit()
        user_id, other_id = user.id, other.id
        token = authenticated_session(session, user).access_token

    def database():
        with Session(engine) as session:
            yield session

    monkeypatch.setattr(get_settings(), "export_dir", str(tmp_path))
    app.dependency_overrides[get_session] = database
    try:
        with TestClient(app) as client:
            client.headers["Authorization"] = f"Bearer {token}"
            yield client, engine, user_id, other_id
    finally:
        app.dependency_overrides.pop(get_session, None)
        engine.dispose()


def test_estimate_survives_request_and_can_generate_proposal(api_client):
    client, engine, _, _ = api_client
    response = client.post(
        "/api/v1/deals/estimate",
        json={"project_type": "dispatching", "scale": "small", "works": []},
    )
    assert response.status_code == 200, response.text
    estimate_id = response.json()["estimate_id"]
    with Session(engine) as session:
        assert session.get(DealEstimate, estimate_id) is not None
    created = client.post("/api/v1/exports/proposal", params={"estimate_id": estimate_id})
    assert created.status_code == 200, created.text
    downloaded = client.get(f"/api/v1/exports/proposal/{estimate_id}/download")
    assert downloaded.status_code == 200
    document = Document(io.BytesIO(downloaded.content))
    assert any("Оценка стоимости" in paragraph.text for paragraph in document.paragraphs)


def _add_analog(engine) -> None:
    # С аналогом в базе расчёт доходит до коэффициента срочности (там и был KeyError).
    with Session(engine) as session:
        session.add(
            HistoricalDeal(
                title="Аналог", project_type="connectivity", scale="medium",
                contract_value=1_000_000,
            )
        )
        session.commit()


def test_estimate_rejects_unknown_urgency(api_client):
    client, engine, _, _ = api_client
    _add_analog(engine)
    response = client.post(
        "/api/v1/deals/estimate",
        json={"project_type": "connectivity", "scale": "medium", "urgency": "срочно"},
    )
    assert response.status_code == 422, response.text


@pytest.mark.parametrize(
    "payload",
    [
        {"project_type": "connectivity", "scale": "огромный"},
        {"project_type": "", "scale": "medium"},
    ],
    ids=["unknown-scale", "empty-project-type"],
)
def test_estimate_rejects_invalid_fields(api_client, payload):
    client, _, _, _ = api_client
    response = client.post("/api/v1/deals/estimate", json=payload)
    assert response.status_code == 422, response.text


def test_estimate_accepts_valid_urgency_with_analogs(api_client):
    client, engine, _, _ = api_client
    _add_analog(engine)
    response = client.post(
        "/api/v1/deals/estimate",
        json={"project_type": "connectivity", "scale": "medium", "urgency": "critical"},
    )
    assert response.status_code == 200, response.text
    assert float(response.json()["estimate"]["optimum"]) > 0


def _seed_calculator(engine) -> None:
    """Таблицы и данные кейса нового калькулятора — для запасного расчёта без аналогов."""
    create_tables(engine)
    with Session(engine) as session:
        load_calculator_bundle(session, json.loads(DEFAULT_BUNDLE.read_text(encoding="utf-8")))
        session.commit()


def _estimate(client, **fields):
    payload = {"project_type": "connectivity", "scale": "medium", **fields}
    response = client.post("/api/v1/deals/estimate", json=payload)
    assert response.status_code == 200, response.text
    return response.json()


def test_estimate_without_analogs_uses_calculator_fallback(api_client):
    # Раньше без исторических аналогов вилка была 0 / 0 / 0 (задача 1).
    client, engine, _, _ = api_client
    _seed_calculator(engine)
    body = _estimate(client)
    assert body["estimator_mode"] == FALLBACK_MODE
    assert 0 < float(body["estimate"]["min"]) <= float(body["estimate"]["optimum"])
    assert float(body["estimate"]["optimum"]) <= float(body["estimate"]["max"])
    assert "5 000 м²" in body["warning"] and "ручная проверка" in body["warning"]
    assert any(a["factor"] == "assumed_area_m2" for a in body["adjustments"])

    critical = _estimate(client, urgency="critical")
    ratio = float(critical["estimate"]["optimum"]) / float(body["estimate"]["optimum"])
    assert ratio == pytest.approx(1.12, rel=1e-3), "срочность учитывается и в запасном расчёте"


def test_fallback_mode_fits_estimator_mode_column():
    # SQLite не проверяет длину строк, PostgreSQL — проверяет: длиннее колонки → 500.
    assert len(FALLBACK_MODE) <= DealEstimate.__table__.c.estimator_mode.type.length


def test_estimate_with_analogs_does_not_use_fallback(api_client):
    client, engine, _, _ = api_client
    _seed_calculator(engine)
    _add_analog(engine)
    body = _estimate(client)
    assert body["estimator_mode"] == "analogs"
    assert not (body["warning"] or "").startswith("Аналогов нет")


def test_fallback_estimate_still_produces_docx(api_client):
    client, engine, _, _ = api_client
    _seed_calculator(engine)
    estimate_id = _estimate(client)["estimate_id"]
    created = client.post("/api/v1/exports/proposal", params={"estimate_id": estimate_id})
    assert created.status_code == 200, created.text
    downloaded = client.get(f"/api/v1/exports/proposal/{estimate_id}/download")
    assert downloaded.status_code == 200
    text = "\n".join(p.text for p in Document(io.BytesIO(downloaded.content)).paragraphs)
    assert "assumed_area_m2: 5000" in text, "допущение видно в КП"
    assert "Минимум: 0 руб." not in text


def test_legacy_estimate_csv_uses_deal_estimate_id(api_client):
    # Раньше CSV запрашивался у /calculator/estimates/{id}/csv с номером из другой таблицы.
    client, engine, _, _ = api_client
    _seed_calculator(engine)
    estimate_id = _estimate(client)["estimate_id"]
    response = client.get(f"/api/v1/deals/estimates/{estimate_id}/csv")
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("text/csv")
    assert response.content.startswith(b"\xef\xbb\xbf"), "UTF-8 BOM для Excel"
    lines = response.content.decode("utf-8-sig").splitlines()
    assert lines[0] == "Раздел,Показатель,Значение,Пояснение"
    assert any(line.startswith("Вилка,Оптимум,") for line in lines)
    assert any(line.startswith("Поправка,assumed_area_m2,5000") for line in lines)


def test_legacy_estimate_csv_is_author_only(api_client):
    client, engine, _, other_id = api_client
    with Session(engine) as session:
        deal = Deal(author_id=other_id, input_params_json={})
        session.add(deal)
        session.flush()
        foreign = DealEstimate(
            deal_id=deal.id, price_min=1, price_optimum=2, price_max=3, model_version="test"
        )
        session.add(foreign)
        session.commit()
        foreign_id = foreign.id
    assert client.get(f"/api/v1/deals/estimates/{foreign_id}/csv").status_code == 404
    assert client.get("/api/v1/deals/estimates/999999/csv").status_code == 404


def test_notifications_have_timestamp_and_remain_user_scoped(api_client):
    client, engine, user_id, other_id = api_client
    with Session(engine) as session:
        own = Notification(
            user_id=user_id,
            type="investment.new",
            payload_json={"title": "Test"},
            dedup_key="own-event",
        )
        other = Notification(
            user_id=other_id, type="investment.new", payload_json={}, dedup_key="other-event"
        )
        session.add_all([own, other])
        session.commit()
        own_id, other_notification_id = own.id, other.id
    response = client.get("/api/v1/notifications")
    assert response.status_code == 200, response.text
    assert [row["id"] for row in response.json()] == [own_id]
    assert response.json()[0]["created_at"]
    assert client.patch(f"/api/v1/notifications/{other_notification_id}/read").status_code == 404
    assert client.patch(f"/api/v1/notifications/{own_id}/read").status_code == 200
    assert client.get("/api/v1/notifications").json() == []
    with Session(engine) as session:
        assert session.scalar(select(Notification.read_at).where(Notification.id == own_id))
