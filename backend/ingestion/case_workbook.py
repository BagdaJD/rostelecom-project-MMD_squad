"""Import the two business sheets of the supplied case, without inventing missing facts."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import UTC, date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import (
    Company,
    CompanyMatch,
    InvestmentEvent,
    InvestmentProject,
    ProjectMilestone,
    SourceRecord,
)
from ingestion.loader import (
    _source,
    load_companies,
    load_dictionaries,
    load_industries,
    load_regions,
)
from ingestion.normalize import RawCompany, is_valid_inn, normalize_name
from modules.investments.service import refresh_needs

COMPANY_SOURCE = "case_msp"
PROJECT_SOURCE = "case_investments"
LICENSE = "Данные предоставленного кейса хакатона; источник — исходный Excel"
REGION = {"code": "RU-ROS-RND", "name": "Ростовская область"}
INDUSTRIES = [
    ("APK", "Агропромышленный комплекс", "АПК"),
    ("TEK", "Топливно-энергетический комплекс", "ТЭК"),
    ("TRADE", "Торговля", "Торговля"),
    ("TRANSPORT", "Транспорт и логистика", "Транспорт"),
    ("IT", "Информационные технологии и связь", "ИТ"),
    ("HEALTH", "Здравоохранение", "Здравоохранение"),
    ("CONSTRUCTION", "Строительство", "Строительство"),
    ("MANUFACTURING", "Промышленность", "Промышленность"),
    ("REAL_ESTATE", "Недвижимость", "Недвижимость"),
    ("TOURISM", "Гостиницы, питание и отдых", "Туризм"),
    ("EDUCATION", "Образование", "Образование"),
    ("FINANCE", "Финансы", "Финансы"),
    ("OTHER", "Другие виды деятельности", "Прочее"),
]


def text(value) -> str:
    return re.sub(r"\s+", " ", str(value)).strip() if value is not None else ""


def industry_for_okved(code: str) -> str:
    prefix = int(code.split(".")[0]) if code and code.split(".")[0].isdigit() else 0
    if prefix in (1, 2, 3, 10, 11, 12):
        return "APK"
    if prefix in (5, 6, 7, 8, 9, 19, 35):
        return "TEK"
    if 45 <= prefix <= 47:
        return "TRADE"
    if 49 <= prefix <= 53:
        return "TRANSPORT"
    if 58 <= prefix <= 63:
        return "IT"
    if 86 <= prefix <= 88:
        return "HEALTH"
    if 41 <= prefix <= 43:
        return "CONSTRUCTION"
    if 13 <= prefix <= 33:
        return "MANUFACTURING"
    return {
        68: "REAL_ESTATE",
        55: "TOURISM",
        56: "TOURISM",
        85: "EDUCATION",
        64: "FINANCE",
        65: "FINANCE",
        66: "FINANCE",
    }.get(prefix, "OTHER")


def municipality(value) -> str | None:
    value = text(value)
    value = re.sub(r"^г\.\s*", "", value, flags=re.I)
    value = re.sub(r"\s+р-н$", " район", value, flags=re.I)
    return value or None


def number(value) -> Decimal | None:
    if value is None or text(value) in ("", "—", "-"):
        return None
    try:
        result = Decimal(text(value).replace(" ", "").replace(",", "."))
    except InvalidOperation as error:
        raise ValueError("Некорректное числовое значение") from error
    if not result.is_finite() or result < 0:
        raise ValueError("Число должно быть конечным и неотрицательным")
    return result


def year(value) -> int | None:
    if value is None or text(value) in ("", "—", "-"):
        return None
    result = int(value)
    if not 1900 <= result <= 2200:
        raise ValueError("Год вне допустимого диапазона")
    return result


def _json_row(row: dict) -> dict:
    return {
        key: value.isoformat() if isinstance(value, (date, datetime)) else value
        for key, value in row.items()
    }


@dataclass
class CaseData:
    companies: list[RawCompany] = field(default_factory=list)
    projects: list[dict] = field(default_factory=list)
    issues: list[dict] = field(default_factory=list)
    company_rows: int = 0
    project_rows: int = 0


def read_case(path: Path) -> CaseData:
    result = CaseData()
    timestamp = datetime.now(UTC)
    signatures = (
        {"Наименование / ФИО", "ИНН", "ОГРН", "Основной вид деятельности"},
        {"Инициатор инвестиционного проекта", "Всего (млн руб.)", "Муниципальное образование"},
    )
    found: set[int] = set()
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        for sheet in workbook:
            iterator = sheet.iter_rows(values_only=True)
            header = [text(value) for value in next(iterator, ())]
            kind = next(
                (i for i, signature in enumerate(signatures) if signature <= set(header)), None
            )
            if kind is None:
                continue
            if kind in found:
                raise ValueError("Неоднозначные листы: схема данных встречается несколько раз")
            found.add(kind)
            seen: set[str] = set()
            for row_number, cells in enumerate(iterator, 2):
                if not any(value is not None for value in cells):
                    continue
                row = dict(zip(header, cells, strict=False))
                if kind == 0:
                    result.company_rows += 1
                else:
                    result.project_rows += 1
                identity = text(row.get("№ п/п"))
                location = {"file": path.name, "sheet": sheet.title, "row": row_number}
                try:
                    if not identity or identity in seen:
                        raise ValueError("Нет уникального номера строки в источнике")
                    seen.add(identity)
                    if kind == 0:
                        result.companies.append(_company(row, identity, location, timestamp))
                    else:
                        result.projects.append(_project(row, identity, location))
                except (ValueError, TypeError) as error:
                    result.issues.append({**location, "external_id": identity, "error": str(error)})
        if found != {0, 1}:
            raise ValueError("В книге не найдены оба листа с компаниями и инвестпроектами")
    finally:
        workbook.close()
    return result


def _company(row: dict, identity: str, location: dict, timestamp: datetime) -> RawCompany:
    name, inn = text(row.get("Наименование / ФИО")), text(row.get("ИНН"))
    if not name or not is_valid_inn(inn):
        raise ValueError("Нет названия или ИНН не прошёл проверку контрольных разрядов")
    if not text(row.get("Регион")).startswith("61"):
        raise ValueError("Для этого импорта поддержан регион 61 из исходного кейса")
    activity = text(row.get("Основной вид деятельности"))
    match = re.match(r"\d{2}(?:\.\d{1,2}){0,2}", activity)
    okved = match.group() if match else ""
    employee_field = next((key for key in row if key.startswith("Среднесписочная")), "")
    employees = number(row.get(employee_field))
    if employees is not None and employees != employees.to_integral_value():
        raise ValueError("Численность работников должна быть целым числом")
    city = text(row.get("Город")) or None
    return RawCompany(
        source=COMPANY_SOURCE,
        external_id=identity,
        name=name,
        inn=inn,
        ogrn=text(row.get("ОГРН")) or None,
        industry_code=okved,
        region_code=REGION["code"],
        employee_count=int(employees) if employees is not None else None,
        municipality=municipality(row.get("Район")) or municipality(city),
        city=city,
        settlement=text(row.get("Населенный пункт")) or None,
        fetched_at=timestamp,
        license_note=LICENSE,
        attributes={
            "provenance": location,
            "raw_row": _json_row(row),
            "okved": okved,
            "industry_method": "okved_prefix",
            "registry_new": text(row.get("Вновь созданный")) == "Да",
        },
    )


def _project(row: dict, identity: str, location: dict) -> dict:
    title_field = next((key for key in row if key.startswith("Название и суть")), "")
    title = text(row.get(title_field))
    initiator = text(row.get("Инициатор инвестиционного проекта"))
    if not title or not initiator:
        raise ValueError("Нет описания проекта или инициатора")
    start, end = year(row.get("Начало")), year(row.get("Ввод в эксплуатацию"))
    if start and end and start > end:
        raise ValueError("Год начала проекта позже года ввода")
    budget = number(row.get("Всего (млн руб.)"))
    stage = "unknown"
    evidence = None
    for marker, candidate in (
        ("модернизац", "modernization"),
        ("реконструкц", "modernization"),
        ("проектирован", "design"),
        ("строительств", "construction"),
    ):
        if marker in title.casefold():
            stage, evidence = candidate, title
            break
    group = "OTHER"
    for markers, candidate in (
        (("сельск", "комбикорм", "животновод", "теплич", "агро"), "APK"),
        (("логист", "склад", "транспорт"), "TRANSPORT"),
        (("торгов", "магазин"), "TRADE"),
        (("завод", "производств", "промышлен"), "MANUFACTURING"),
        (("гостиниц", "отдых", "спортив"), "TOURISM"),
    ):
        if any(marker in title.casefold() for marker in markers):
            group = candidate
            break
    return {
        "external_id": identity,
        "title": title,
        "initiator": initiator,
        "municipality": municipality(row.get("Муниципальное образование")),
        "start_year": start,
        "end_year": end,
        "budget_rub": str(budget * 1_000_000) if budget is not None else None,
        "stage": stage,
        "stage_evidence": evidence,
        "industry": group,
        "provenance": location,
        "raw_row": _json_row(row),
    }


def import_case(session: Session, path: Path) -> dict:
    data = read_case(path)
    load_dictionaries(session)
    region_id = load_regions(session, [REGION])[REGION["code"]]
    rows = [
        {
            "code": code,
            "name": name,
            "group": group,
            "okved": sorted(
                {
                    c.industry_code
                    for c in data.companies
                    if industry_for_okved(c.industry_code or "") == code
                }
            ),
        }
        for code, name, group in INDUSTRIES
    ]
    industries, okved = load_industries(session, rows)
    _, companies = load_companies(session, data.companies, okved, {REGION["code"]: region_id})

    session.flush()
    created = updated = initiators = 0
    for row in data.projects:
        source = _source(session, PROJECT_SOURCE, row["external_id"], row, LICENSE)
        project = session.scalar(
            select(InvestmentProject)
            .join(SourceRecord)
            .where(
                SourceRecord.source == PROJECT_SOURCE,
                SourceRecord.external_id == row["external_id"],
            )
        )
        changed = project is None or project.source_id != source.id
        is_new = project is None
        normalized = normalize_name(row["initiator"])
        candidates = session.scalars(
            select(Company)
            .where(Company.name_normalized == normalized, Company.region_id == region_id)
            .limit(2)
        ).all()
        initiator = candidates[0] if len(candidates) == 1 else None
        if initiator is None and len(candidates) == 0:
            initiator = Company(
                name=row["initiator"],
                name_normalized=normalized,
                region_id=region_id,
                source_id=source.id,
                industry_id=industries[row["industry"]],
                municipality=row["municipality"],
                client_status="prospect",
                first_seen_at=source.fetched_at,
                last_seen_at=source.fetched_at,
                attributes_json={"observed_as": "project_initiator"},
            )
            session.add(initiator)
            session.flush()
            initiators += 1
        if initiator and initiator.client_status != "active":
            initiator.client_status = "prospect"
        if initiator and not session.scalar(
            select(CompanyMatch.id).where(
                CompanyMatch.company_id == initiator.id, CompanyMatch.source_record_id == source.id
            )
        ):
            session.add(
                CompanyMatch(
                    company_id=initiator.id,
                    source_record_id=source.id,
                    match_method="name_region",
                    confidence=Decimal("0.750"),
                )
            )
        if project is None:
            project = InvestmentProject(title=row["title"])
            session.add(project)
            created += 1
        elif changed:
            updated += 1
        project.title = row["title"]
        project.description = row["title"]
        project.initiator_company_id = initiator.id if initiator else None
        project.region_id = region_id
        project.industry_id = industries[row["industry"]]
        project.municipality = row["municipality"]
        project.source_id = source.id
        project.stage = row["stage"]
        project.stage_method = "rule_based" if row["stage_evidence"] else "unknown"
        project.stage_evidence = row["stage_evidence"]
        project.confidence = Decimal("0.500") if row["stage_evidence"] else None
        project.planned_start = date(row["start_year"], 1, 1) if row["start_year"] else None
        project.planned_end = date(row["end_year"], 12, 31) if row["end_year"] else None
        project.start_precision = "year" if row["start_year"] else "unknown"
        project.end_precision = "year" if row["end_year"] else "unknown"
        project.estimated_budget = (
            Decimal(row["budget_rub"]) if row["budget_rub"] is not None else None
        )
        project.attributes_json = {
            "provenance": row["provenance"],
            "industry_method": "rule_based",
            "initiator_match": "name_region" if initiator else "ambiguous",
        }
        session.flush()
        if changed:
            refresh_needs(session, project, "needs-rules-0.1")
            from modules.notifications.events import project_event

            project_event(session, project, created=is_new)
        for label, value in (
            ("planned_start", project.planned_start),
            ("planned_end", project.planned_end),
        ):
            event = session.scalar(
                select(InvestmentEvent).where(
                    InvestmentEvent.project_id == project.id, InvestmentEvent.event_type == label
                )
            )
            milestone = session.scalar(
                select(ProjectMilestone).where(
                    ProjectMilestone.project_id == project.id, ProjectMilestone.stage == label
                )
            )
            if value is None:
                if event:
                    session.delete(event)
                if milestone:
                    session.delete(milestone)
                continue
            if event is None:
                event = InvestmentEvent(project_id=project.id, event_type=label)
                session.add(event)
            event.event_date, event.date_precision = value, "year"
            event.region_id, event.industry_id, event.title = (
                region_id,
                project.industry_id,
                project.title,
            )
            if milestone is None:
                milestone = ProjectMilestone(project_id=project.id, stage=label)
                session.add(milestone)
            milestone.planned_date, milestone.date_precision = value, "year"
        session.flush()
    from modules.radar.signals import detect_signals
    from modules.readmodels.service import mark_changed

    detect_signals(session)
    mark_changed(session)
    return {
        "company_rows": data.company_rows,
        "project_rows": data.project_rows,
        "companies_created": companies.companies,
        "initiators_created": initiators,
        "projects_created": created,
        "projects_updated": updated,
        "accepted_companies": len(data.companies),
        "accepted_projects": len(data.projects),
        "rejected_rows": len(data.issues),
        "issues": data.issues,
    }
